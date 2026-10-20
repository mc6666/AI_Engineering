import csv
import math
from pathlib import Path
import tkinter as tk
from dataclasses import dataclass
from tkinter import filedialog, messagebox, simpledialog, ttk

from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure
from matplotlib.patches import FancyArrowPatch
from collections import defaultdict
from loguru import logger

@dataclass
class PendingAction:
    kind: str
    text: str = ""
    start: tuple[float, float] | None = None


@dataclass
class DragState:
    kind: str
    axes: object
    artist: object
    start_mouse: tuple[float, float]
    start_data: tuple[float, float]
    payload: dict


class MatplotlibDesignerApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Matplotlib Designer")
        self.root.geometry("1440x900")
        self.root.minsize(1200, 760)

        self.pending_action: PendingAction | None = None
        self.drag_state: DragState | None = None
        self.axes_items: list[dict] = []
        self.active_axes_index: int = 0
        self.figure_title = ""
        self.figure_subtitle = ""
        self.subtitle_artist = None
        self.data_source_path: str = ""
        self.data_columns: list[str] = []
        self.data_rows: list[dict[str, object]] = []

        self._build_ui()
        self._create_initial_axes()
        self._bind_canvas_events()
        self.refresh_all()

    def _build_ui(self) -> None:
        self.root.configure(bg="#1f232a")

        self.toolbar = ttk.Frame(self.root, padding=(10, 8))
        self.toolbar.pack(side=tk.TOP, fill=tk.X)

        self._add_toolbar_button("Load Data", self.load_data_dialog)
        self._add_toolbar_button("Plot Data", self.plot_loaded_data)
        self._add_toolbar_button("Reset", self.reset_figure)
        self._add_toolbar_button("Title", self.set_title)
        self._add_toolbar_button("Subtitle", self.set_subtitle)
        self._add_toolbar_button("X Label", self.set_xlabel)
        self._add_toolbar_button("Y Label", self.set_ylabel)
        self._add_toolbar_button("Ticks", self.set_ticks)
        self._add_toolbar_button("Annotation", self.add_annotation_mode)
        self._add_toolbar_button("Arrow + Text", self.add_arrow_mode)
        self._add_toolbar_button("Add Subplot", self.add_subplot_dialog)
        # self._add_toolbar_button("Sample Plot", self.add_sample_plot)
        self._add_toolbar_button("Export to Code", self.export_to_code)
        self._add_toolbar_button("Save PNG", self.save_png)

        self.main = ttk.Frame(self.root, padding=(10, 0, 10, 10))
        self.main.pack(side=tk.TOP, fill=tk.BOTH, expand=True)

        self.left_panel = ttk.Frame(self.main)
        self.left_panel.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self.right_panel = ttk.Frame(self.main, width=280)
        self.right_panel.pack(side=tk.RIGHT, fill=tk.Y, padx=(10, 0))
        self.right_panel.pack_propagate(False)

        self.figure = Figure(figsize=(8, 6), dpi=100)
        self.canvas = FigureCanvasTkAgg(self.figure, master=self.left_panel)
        self.canvas_widget = self.canvas.get_tk_widget()
        self.canvas_widget.pack(side=tk.TOP, fill=tk.BOTH, expand=True)

        sidebar_title = ttk.Label(self.right_panel, text="Figure Controls", font=("Segoe UI", 13, "bold"))
        sidebar_title.pack(anchor="w", pady=(0, 10))

        axes_box = ttk.LabelFrame(self.right_panel, text="Active Axes", padding=10)
        axes_box.pack(fill=tk.X, pady=(0, 12))

        self.axes_list = tk.Listbox(axes_box, height=7, exportselection=False)
        self.axes_list.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        axes_scroll = ttk.Scrollbar(axes_box, orient=tk.VERTICAL, command=self.axes_list.yview)
        axes_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.axes_list.configure(yscrollcommand=axes_scroll.set)
        self.axes_list.bind("<<ListboxSelect>>", self._on_axes_selected)

        ttk.Label(
            self.right_panel,
            text=(
                "Click the canvas to place annotations and arrows.\n"
                "Arrow mode uses two clicks.\n"
                "Drag annotation boxes and arrow labels after placing them."
            ),
            justify=tk.LEFT,
        ).pack(anchor="w", pady=(0, 12))

        data_box = ttk.LabelFrame(self.right_panel, text="Loaded Data", padding=10)
        data_box.pack(fill=tk.BOTH, expand=False, pady=(0, 12))

        self.data_summary = ttk.Label(data_box, text="No file loaded.", justify=tk.LEFT, wraplength=240)
        self.data_summary.pack(anchor="w", pady=(0, 8))

        ttk.Label(data_box, text="Plot type").pack(anchor="w")
        self.plot_type_combo = ttk.Combobox(data_box, state="readonly", values=["count", "line", "scatter", "bar", "hist"])
        self.plot_type_combo.pack(fill=tk.X, pady=(0, 8))
        self.plot_type_combo.set("count")

        ttk.Label(data_box, text="X field").pack(anchor="w")
        self.x_field_combo = ttk.Combobox(data_box, state="readonly", values=[])
        self.x_field_combo.pack(fill=tk.X, pady=(0, 8))

        ttk.Label(data_box, text="Y fields").pack(anchor="w")
        y_frame = ttk.Frame(data_box)
        y_frame.pack(fill=tk.BOTH, expand=True)
        self.y_field_listbox = tk.Listbox(y_frame, height=6, selectmode=tk.EXTENDED, exportselection=False)
        self.y_field_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        y_scroll = ttk.Scrollbar(y_frame, orient=tk.VERTICAL, command=self.y_field_listbox.yview)
        y_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.y_field_listbox.configure(yscrollcommand=y_scroll.set)

        ttk.Button(data_box, text="Plot Selected Fields", command=self.plot_loaded_data).pack(fill=tk.X, pady=(8, 0))

        # quick_box = ttk.LabelFrame(self.right_panel, text="Quick Actions", padding=10)
        # quick_box.pack(fill=tk.X, pady=(0, 12))
        # ttk.Button(quick_box, text="Add Subplot", command=self.add_subplot_dialog).pack(fill=tk.X, pady=3)
        # ttk.Button(quick_box, text="Add Sample Plot", command=self.add_sample_plot).pack(fill=tk.X, pady=3)
        # ttk.Button(quick_box, text="Load Data", command=self.load_data_dialog).pack(fill=tk.X, pady=3)
        # ttk.Button(quick_box, text="Plot Data", command=self.plot_loaded_data).pack(fill=tk.X, pady=3)
        # ttk.Button(quick_box, text="Export to Code", command=self.export_to_code).pack(fill=tk.X, pady=3)
        # ttk.Button(quick_box, text="Reset Figure", command=self.reset_figure).pack(fill=tk.X, pady=3)
        # ttk.Button(quick_box, text="Save PNG", command=self.save_png).pack(fill=tk.X, pady=3)

        self.status = tk.StringVar(value="Ready.")
        status_bar = ttk.Label(self.root, textvariable=self.status, anchor="w", padding=(10, 4))
        status_bar.pack(side=tk.BOTTOM, fill=tk.X)

    def _add_toolbar_button(self, label: str, command) -> None:
        ttk.Button(self.toolbar, text=label, command=command).pack(side=tk.LEFT, padx=4)

    def _create_initial_axes(self) -> None:
        ax = self.figure.add_subplot(1, 1, 1)
        self.axes_items.append({"ax": ax, "label": "Axes 1", "spec": (1, 1, 1)})
        self.axes_list.insert(tk.END, "Axes 1")
        self.axes_list.selection_set(0)
        self.axes_list.activate(0)
        self.active_axes_index = 0
        # self._draw_demo(ax)

    def _bind_canvas_events(self) -> None:
        self.canvas.mpl_connect("button_press_event", self.on_canvas_click)
        self.canvas.mpl_connect("pick_event", self.on_pick)
        self.canvas.mpl_connect("motion_notify_event", self.on_motion)
        self.canvas.mpl_connect("button_release_event", self.on_release)

    def active_axes(self):
        if not self.axes_items:
            return None
        index = max(0, min(self.active_axes_index, len(self.axes_items) - 1))
        return self.axes_items[index]["ax"]

    def _set_status(self, text: str) -> None:
        self.status.set(text)

    def _refresh_axes_list(self) -> None:
        self.axes_list.delete(0, tk.END)
        for idx, _item in enumerate(self.axes_items, start=1):
            self.axes_list.insert(tk.END, f"Axes {idx}")
        if self.axes_items:
            self.active_axes_index = max(0, min(self.active_axes_index, len(self.axes_items) - 1))
            self.axes_list.selection_clear(0, tk.END)
            self.axes_list.selection_set(self.active_axes_index)
            self.axes_list.activate(self.active_axes_index)

    def _on_axes_selected(self, _event=None) -> None:
        selection = self.axes_list.curselection()
        if not selection:
            return
        self.active_axes_index = selection[0]
        self._set_status(f"Selected Axes {self.active_axes_index + 1}.")

    def _draw_demo(self, ax) -> None:
        xs = [i / 25 for i in range(0, 251)]
        ys = [math.sin(x) for x in xs]
        ax.plot(xs, ys, color="#1f77b4", linewidth=2)
        ax.set_title("Sample Plot", fontsize=11)
        ax.set_xlabel("x")
        ax.set_ylabel("sin(x)")
        ax._designer_custom_xticks = False
        ax._designer_custom_yticks = False
        self._apply_pickers(ax)

    def _apply_pickers(self, ax) -> None:
        for text in ax.texts:
            text.set_picker(True)
        for artist in ax.patches:
            artist.set_picker(5)

    def _prompt_text(self, title: str, prompt: str, default: str = "") -> str | None:
        value = simpledialog.askstring(title, prompt, initialvalue=default, parent=self.root)
        if value is None:
            return None
        value = value.strip()
        return value if value else None

    def _prompt_float_list(self, title: str, prompt: str) -> list[float] | None:
        raw = simpledialog.askstring(title, prompt, parent=self.root)
        if raw is None:
            return None
        raw = raw.strip()
        if not raw:
            return []
        values: list[float] = []
        for chunk in raw.split(","):
            text = chunk.strip()
            if not text:
                continue
            try:
                values.append(float(text))
            except ValueError:
                messagebox.showerror("Invalid value", f"Could not parse '{text}' as a number.")
                return None
        return values

    def _coerce_number(self, value: object) -> float | None:
        if value is None:
            return None
        if isinstance(value, bool):
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip()
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None

    def _normalize_headers(self, headers: list[object]) -> list[str]:
        result: list[str] = []
        seen: set[str] = set()
        for index, header in enumerate(headers, start=1):
            name = str(header).strip() if header is not None else ''
            if not name:
                name = f'Column {index}'
            base_name = name
            suffix = 2
            while name in seen:
                name = f'{base_name} ({suffix})'
                suffix += 1
            seen.add(name)
            result.append(name)
        return result

    def _read_csv_table(self, path: Path) -> tuple[list[str], list[dict[str, object]]]:
        with path.open('r', encoding='utf-8-sig', newline='') as handle:
            sample = handle.read(4096)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',	;|')
            except csv.Error:
                dialect = csv.get_dialect('excel')
            reader = csv.reader(handle, dialect)
            rows = list(reader)
        if not rows:
            raise ValueError('The file is empty.')
        headers = self._normalize_headers(rows[0])
        data_rows: list[dict[str, object]] = []
        for raw_row in rows[1:]:
            if not any(str(cell).strip() for cell in raw_row):
                continue
            record: dict[str, object] = {}
            for index, header in enumerate(headers):
                record[header] = raw_row[index] if index < len(raw_row) else ''
            data_rows.append(record)
        if not data_rows:
            raise ValueError('The file does not contain any data rows.')
        return headers, data_rows

    def _read_excel_table(self, path: Path) -> tuple[list[str], list[dict[str, object]]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel support requires the 'openpyxl' package.") from exc

        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        if not rows:
            raise ValueError('The workbook is empty.')
        headers = self._normalize_headers(list(rows[0]))
        data_rows: list[dict[str, object]] = []
        for raw_row in rows[1:]:
            if not any(cell is not None and str(cell).strip() for cell in raw_row):
                continue
            record: dict[str, object] = {}
            for index, header in enumerate(headers):
                record[header] = raw_row[index] if index < len(raw_row) else None
            data_rows.append(record)
        if not data_rows:
            raise ValueError('The workbook does not contain any data rows.')
        return headers, data_rows

    def _read_tabular_file(self, path: Path) -> tuple[list[str], list[dict[str, object]]]:
        suffix = path.suffix.lower()
        if suffix == '.csv':
            return self._read_csv_table(path)
        if suffix in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
            return self._read_excel_table(path)
        raise ValueError('Unsupported file type. Use CSV or Excel (.xlsx, .xlsm, .xltx, .xltm).')

    def _update_data_controls(self) -> None:
        self.x_field_combo.configure(values=self.data_columns)
        self.y_field_listbox.delete(0, tk.END)

        if not self.data_columns:
            self.x_field_combo.set('')
            self.data_summary.configure(text='No file loaded.')
            return

        default_x = self.data_columns[0]
        default_ys = self.data_columns[1:] if len(self.data_columns) > 1 else []
        self.x_field_combo.set(default_x)
        if self.plot_type_combo.get() not in {"count", "bar", "line", "scatter"}:
            self.plot_type_combo.set("count")
        for column in self.data_columns:
            self.y_field_listbox.insert(tk.END, column)
        for index, column in enumerate(self.data_columns):
            if column in default_ys:
                self.y_field_listbox.selection_set(index)

        summary = f"File: {Path(self.data_source_path).name}\nColumns: {len(self.data_columns)}\nRows: {len(self.data_rows)}"
        self.data_summary.configure(text=summary)

    def _selected_y_fields(self) -> list[str]:
        selection = self.y_field_listbox.curselection()
        if not selection:
            return []
        return [self.data_columns[index] for index in selection if 0 <= index < len(self.data_columns)]

    def load_data_dialog(self) -> None:
        path = filedialog.askopenfilename(
            parent=self.root,
            title='Load CSV or Excel File',
            filetypes=[
                ('CSV / Excel', '*.csv *.xlsx *.xlsm *.xltx *.xltm'),
                ('CSV', '*.csv'),
                ('Excel', '*.xlsx *.xlsm *.xltx *.xltm'),
                ('All Files', '*.*'),
            ],
        )
        if not path:
            return

        try:
            columns, rows = self._read_tabular_file(Path(path))
        except Exception as exc:
            messagebox.showerror('Load Data', str(exc))
            return

        self.data_source_path = path
        self.data_columns = columns
        self.data_rows = rows
        self._update_data_controls()
        self._set_status(f'Loaded {Path(path).name} with {len(rows)} rows.')

    def _clear_data_plot_artists(self, ax) -> None:
        # legend = ax.get_legend()
        # if legend is not None:
            # legend.remove()
        # existing = getattr(ax, '_designer_data_artists', [])
        # for artist in existing:
            # try:
                # artist.remove()
            # except Exception:
                # pass
        # ax._designer_data_artists = []
        # ax._designer_data_plots = []
        ax.cla()
        
    def group_data(self, x, y, aggregation_type = 'average'):
        # Step 1: Group values by category
        grouped_data = defaultdict(list)
        for cat, val in zip(x, y):
            grouped_data[cat].append(val)

        # Step 2: Compute aggregation into a dictionary
        if aggregation_type == 'average':
            group_dict = {cat: sum(vals) / len(vals) for cat, vals in grouped_data.items()}
        elif aggregation_type == 'sum':
            group_dict = {cat: sum(vals) for cat, vals in grouped_data.items()}
        else:
            group_dict = {cat: len(vals) for cat, vals in grouped_data.items()}
        logger.debug(group_dict)
        return group_dict.keys(), group_dict.values()
        
    def plot_loaded_data(self) -> None:
        logger.debug('plot_loaded_data.')
        ax = self.active_axes()
        if ax is None:
            return
        if not self.data_rows or not self.data_columns:
            messagebox.showwarning('Plot Data', 'Load a CSV or Excel file first.')
            return

        x_field = self.x_field_combo.get().strip()
        if not x_field:
            messagebox.showwarning('Plot Data', 'Select an X field.')
            return
        if x_field not in self.data_columns:
            messagebox.showwarning('Plot Data', 'Select a valid X field from the loaded data.')
            return

        plot_type = self.plot_type_combo.get().strip().lower()
        # if plot_type not in {"count", "line", "scatter", "bar", "hist"}:
            # plot_type = "line"

        y_fields = self._selected_y_fields()
        if not y_fields:
            y_fields = [field for field in self.data_columns if field != x_field]
        if plot_type != "hist":
            y_fields = [field for field in y_fields if field != x_field]
        if not y_fields:
            messagebox.showwarning('Plot Data', 'Select at least one Y field.')
            return

        self._clear_data_plot_artists(ax)

        plotted_fields: list[str] = []
        data_plot_records: list[dict[str, object]] = []

        if plot_type == "hist":
            for field in y_fields:
                values: list[float] = []
                for row in self.data_rows:
                    value = self._coerce_number(row.get(field))
                    if value is None:
                        continue
                    values.append(value)
                if not values:
                    continue
                hist = ax.hist(values, bins="auto", alpha=0.55, label=field)
                patches = list(hist[2])
                if patches:
                    if not hasattr(ax, '_designer_data_artists'):
                        ax._designer_data_artists = []
                    ax._designer_data_artists.extend(patches)
                color = None
                if patches:
                    color = self._rgba_tuple(patches[0].get_facecolor())
                data_plot_records.append({
                    "kind": "hist",
                    "field": field,
                    "values": values,
                    "bins": "auto",
                    "alpha": 0.55,
                    "color": color,
                })
                plotted_fields.append(field)
        elif plot_type == "bar":
            bar_width = 0.8 / max(1, len(y_fields))
            start_offset = -0.4 + (bar_width / 2)
            for field_index, field in enumerate(y_fields):
                xs = []
                ys: list[float] = []
                for row in self.data_rows:
                    x_value = row.get(x_field)
                    y_value = self._coerce_number(row.get(field))
                    if x_value is None or y_value is None:
                        continue
                    xs.append(x_value)
                    ys.append(y_value)
                if not xs:
                    continue
                
                xs_group = set(xs)
                if len(xs) != len(xs_group):
                    xs, ys = self.group_data(xs, ys)
                        
                offset = start_offset + field_index * bar_width
                # bar_xs = [x + offset for x in xs]
                container = ax.bar(xs, ys, width=bar_width, label=field, align="center")
                patches = list(container.patches)
                if patches:
                    if not hasattr(ax, '_designer_data_artists'):
                        ax._designer_data_artists = []
                    ax._designer_data_artists.extend(patches)
                color = None
                if patches:
                    color = self._rgba_tuple(patches[0].get_facecolor())
                data_plot_records.append({
                    "kind": "bar",
                    "field": field,
                    "x": xs,
                    "y": ys,
                    "width": bar_width,
                    "color": color,
                })
                plotted_fields.append(field)
        elif plot_type == "count":
            logger.debug('select count.')
            bar_width = 0.8 / max(1, len(y_fields))
            start_offset = -0.4 + (bar_width / 2)
            for field_index, field in enumerate(y_fields):
                xs = []
                ys: list[float] = []
                for i, row in enumerate(self.data_rows):
                    x_value = row.get(x_field)
                    y_value = self._coerce_number(row.get(field))
                    if x_value is None or y_value is None:
                        logger.debug(f'row {i} has None. ({x_value}, {y_value})')
                        continue
                    xs.append(x_value)
                    ys.append(y_value)
                if not xs:
                    logger.debug('no rows.')
                    continue
                
                xs_group = set(xs)
                logger.debug(f'{len(xs)} == {len(xs_group)}?')
                if len(xs) != len(xs_group):
                    xs, ys = self.group_data(xs, ys, 'count')
                        
                offset = start_offset + field_index * bar_width
                # bar_xs = [x + offset for x in xs]
                container = ax.bar(xs, ys, width=bar_width, label=field, align="center")
                patches = list(container.patches)
                if patches:
                    if not hasattr(ax, '_designer_data_artists'):
                        ax._designer_data_artists = []
                    ax._designer_data_artists.extend(patches)
                color = None
                if patches:
                    color = self._rgba_tuple(patches[0].get_facecolor())
                data_plot_records.append({
                    "kind": "bar",
                    "field": field,
                    "x": xs,
                    "y": ys,
                    "width": bar_width,
                    "color": color,
                })
                plotted_fields.append(field)
        else:
            for field in y_fields:
                xs: list[float] = []
                ys: list[float] = []
                for row in self.data_rows:
                    x_value = self._coerce_number(row.get(x_field))
                    y_value = self._coerce_number(row.get(field))
                    if x_value is None or y_value is None:
                        continue
                    xs.append(x_value)
                    ys.append(y_value)
                if not xs:
                    continue
                if plot_type == "scatter":
                    artist = ax.scatter(xs, ys, s=24, label=field)
                else:
                    artist = ax.plot(xs, ys, linewidth=2.0, label=field)[0]
                if not hasattr(ax, '_designer_data_artists'):
                    ax._designer_data_artists = []
                ax._designer_data_artists.append(artist)
                color = None
                if plot_type == "scatter":
                    facecolors = artist.get_facecolors()
                    if len(facecolors) > 0:
                        color = self._rgba_tuple(facecolors[0])
                else:
                    color = artist.get_color()
                data_plot_records.append({
                    "kind": plot_type,
                    "field": field,
                    "x": xs,
                    "y": ys,
                    "color": color,
                })
                plotted_fields.append(field)

        ax._designer_data_plots = data_plot_records

        if not plotted_fields:
            messagebox.showwarning(
                'Plot Data',
                'No plottable numeric values were found for the selected fields.',
            )
            return

        if plot_type != "hist":
            ax.set_xlabel(x_field)
        else:
            ax.set_xlabel("Value")
        ax.set_ylabel(', '.join(plotted_fields))
        ax.set_title(Path(self.data_source_path).name, fontsize=13, pad=10)
        ax.grid(True, alpha=0.25)
        if len(plotted_fields) > 1:
            ax.legend()
        self._apply_pickers(ax)
        self._set_status(
            f'Plotted {len(plotted_fields)} field(s) as {plot_type} from {Path(self.data_source_path).name} on Axes {self.active_axes_index + 1}.'
        )
        self.refresh_all()

    def refresh_layout(self) -> None:
        has_header = bool(self.figure_title or self.figure_subtitle)
        top_margin = 0.86 if has_header else 0.93
        self.figure.subplots_adjust(left=0.08, right=0.97, bottom=0.08, top=top_margin, wspace=0.32, hspace=0.4)

    def refresh_header(self) -> None:
        if self.subtitle_artist is not None:
            self.subtitle_artist.remove()
            self.subtitle_artist = None

        if self.figure_title:
            self.figure.suptitle(self.figure_title, fontsize=16, fontweight="bold", y=0.98)
        else:
            self.figure.suptitle("")

        if self.figure_subtitle:
            self.subtitle_artist = self.figure.text(
                0.5,
                0.945,
                self.figure_subtitle,
                ha="center",
                va="top",
                fontsize=10.5,
                color="#4c5563",
            )

    def refresh_all(self) -> None:
        for item in self.axes_items:
            self._apply_pickers(item["ax"])
        self.refresh_header()
        self.refresh_layout()
        self.canvas.draw_idle()

    def set_title(self) -> None:
        ax = self.active_axes()
        if ax is None:
            return
        text = self._prompt_text("Set Title", "Title text:")
        if text is None:
            return
        ax.set_title(text, fontsize=13, pad=10)
        self._apply_pickers(ax)
        self._set_status(f"Title set on Axes {self.active_axes_index + 1}.")
        self.refresh_all()

    def set_subtitle(self) -> None:
        text = self._prompt_text("Set Subtitle", "Figure subtitle text:")
        if text is None:
            return
        self.figure_subtitle = text
        self._set_status("Figure subtitle updated.")
        self.refresh_all()

    def set_xlabel(self) -> None:
        ax = self.active_axes()
        if ax is None:
            return
        text = self._prompt_text("Set X Label", "X label text:")
        if text is None:
            return
        ax.set_xlabel(text)
        self._apply_pickers(ax)
        self._set_status(f"X label set on Axes {self.active_axes_index + 1}.")
        self.refresh_all()

    def set_ylabel(self) -> None:
        ax = self.active_axes()
        if ax is None:
            return
        text = self._prompt_text("Set Y Label", "Y label text:")
        if text is None:
            return
        ax.set_ylabel(text)
        self._apply_pickers(ax)
        self._set_status(f"Y label set on Axes {self.active_axes_index + 1}.")
        self.refresh_all()

    def set_ticks(self) -> None:
        ax = self.active_axes()
        if ax is None:
            return
        x_ticks = self._prompt_float_list("Set X Ticks", "X tick positions, comma-separated:")
        if x_ticks is None:
            return
        y_ticks = self._prompt_float_list("Set Y Ticks", "Y tick positions, comma-separated:")
        if y_ticks is None:
            return

        if x_ticks:
            ax.set_xticks(x_ticks)
            ax._designer_custom_xticks = True
        if y_ticks:
            ax.set_yticks(y_ticks)
            ax._designer_custom_yticks = True

        self._set_status(f"Ticks updated on Axes {self.active_axes_index + 1}.")
        self.refresh_all()

    def add_annotation_mode(self) -> None:
        text = self._prompt_text("Add Annotation", "Annotation text:")
        if text is None:
            return
        self.pending_action = PendingAction(kind="annotation", text=text)
        self._set_status("Annotation mode active. Click on the target point.")

    def add_arrow_mode(self) -> None:
        text = self._prompt_text("Add Arrow + Text", "Arrow label text:")
        if text is None:
            return
        self.pending_action = PendingAction(kind="arrow", text=text)
        self._set_status("Arrow mode active. Click start point, then click end point.")

    def add_subplot_dialog(self) -> None:
        rows = simpledialog.askinteger("Add Subplot", "Number of rows:", minvalue=1, initialvalue=1, parent=self.root)
        if rows is None:
            return
        cols = simpledialog.askinteger(
            "Add Subplot",
            "Number of columns:",
            minvalue=1,
            initialvalue=max(1, len(self.axes_items) + 1),
            parent=self.root,
        )
        if cols is None:
            return
        index = simpledialog.askinteger(
            "Add Subplot",
            "Subplot position index (1-based):",
            minvalue=1,
            maxvalue=rows * cols,
            initialvalue=min(rows * cols, len(self.axes_items) + 1),
            parent=self.root,
        )
        if index is None:
            return

        ax = self.figure.add_subplot(rows, cols, index)
        ax._designer_custom_xticks = False
        ax._designer_custom_yticks = False
        self.axes_items.append({"ax": ax, "label": f"Axes {len(self.axes_items) + 1}", "spec": (rows, cols, index)})
        self._refresh_axes_list()
        self.active_axes_index = len(self.axes_items) - 1
        self.axes_list.selection_clear(0, tk.END)
        self.axes_list.selection_set(self.active_axes_index)
        self.axes_list.activate(self.active_axes_index)
        # self._draw_demo(ax)
        self._set_status(f"Added subplot at {rows}x{cols}, position {index}.")
        self.refresh_all()

    def add_sample_plot(self) -> None:
        ax = self.active_axes()
        if ax is None:
            return
        xs = [i / 20 for i in range(0, 201)]
        ys = [math.sin(x) + 0.35 * math.cos(3 * x) for x in xs]
        ax.plot(xs, ys, linewidth=2.2, color="#d62728")
        ax.scatter(xs[::20], ys[::20], s=18, color="#2ca02c")
        self._apply_pickers(ax)
        self._set_status(f"Sample plot added to Axes {self.active_axes_index + 1}.")
        self.refresh_all()

    def reset_figure(self) -> None:
        if not messagebox.askyesno("Reset Figure", "Clear the current figure and start over?"):
            return
        self.figure.clear()
        self.axes_items.clear()
        self.pending_action = None
        self.drag_state = None
        self.figure_title = ""
        self.figure_subtitle = ""
        self.subtitle_artist = None
        self.data_source_path = ""
        self.data_columns = []
        self.data_rows = []
        self.x_field_combo.configure(values=[])
        self.x_field_combo.set("")
        self.y_field_listbox.delete(0, tk.END)
        self.data_summary.configure(text="No file loaded.")
        self._refresh_axes_list()
        self._create_initial_axes()
        self._set_status("Figure reset.")
        self.refresh_all()

    def save_png(self) -> None:
        path = filedialog.asksaveasfilename(
            parent=self.root,
            title="Save Figure",
            defaultextension=".png",
            filetypes=[("PNG Image", "*.png"), ("All Files", "*.*")],
        )
        if not path:
            return
        self.figure.savefig(path, dpi=200, bbox_inches="tight")
        self._set_status(f"Saved figure to {path}.")

    def export_to_code(self) -> None:
        path = filedialog.asksaveasfilename(
            parent=self.root,
            title="Export Matplotlib Script",
            defaultextension=".py",
            filetypes=[("Python Script", "*.py"), ("All Files", "*.*")],
        )
        if not path:
            return

        script = self._generate_export_script()
        with open(path, "w", encoding="utf-8") as handle:
            handle.write(script)

        self._set_status(f"Exported Matplotlib script to {path}.")
        messagebox.showinfo("Export complete", f"Saved script to:\n{path}")

    def _format_float_list(self, values) -> str:
        return "[" + ", ".join(repr(float(value)) for value in values) + "]"

    def _rgba_tuple(self, value) -> tuple[float, float, float, float] | None:
        if value is None:
            return None
        try:
            return tuple(float(component) for component in value)
        except Exception:
            return None

    def _generate_export_script(self) -> str:
        lines: list[str] = []

        def add(text: str = "") -> None:
            lines.append(text)

        add("import matplotlib.pyplot as plt")
        add("from matplotlib.patches import FancyArrowPatch")
        add("")
        size = tuple(float(value) for value in self.figure.get_size_inches())
        add(f"fig = plt.figure(figsize={size!r}, dpi={int(self.figure.get_dpi())})")

        if self.figure_title:
            add(f"fig.suptitle({self.figure_title!r}, fontsize=16, fontweight='bold', y=0.98)")
        if self.figure_subtitle:
            add(
                f"fig.text(0.5, 0.945, {self.figure_subtitle!r}, ha='center', va='top', "
                "fontsize=10.5, color='#4c5563')"
            )

        has_header = bool(self.figure_title or self.figure_subtitle)
        top_margin = 0.86 if has_header else 0.93
        add(
            "fig.subplots_adjust(left=0.08, right=0.97, bottom=0.08, top="
            f"{top_margin}, wspace=0.32, hspace=0.4)"
        )
        add("")

        arrow_label_ids: set[int] = set()
        for item in self.axes_items:
            for record in self._get_arrow_records(item["ax"]):
                arrow_label_ids.add(id(record["label"]))

        for idx, item in enumerate(self.axes_items, start=1):
            rows, cols, index = item["spec"]
            ax = item["ax"]
            add(f"ax{idx} = fig.add_subplot({rows}, {cols}, {index})")

            title = ax.get_title()
            if title:
                add(f"ax{idx}.set_title({title!r}, fontsize=13, pad=10)")

            xlabel = ax.get_xlabel()
            if xlabel:
                add(f"ax{idx}.set_xlabel({xlabel!r})")

            ylabel = ax.get_ylabel()
            if ylabel:
                add(f"ax{idx}.set_ylabel({ylabel!r})")

            if getattr(ax, "_designer_custom_xticks", False):
                add(f"ax{idx}.set_xticks({self._format_float_list(ax.get_xticks())})")
            if getattr(ax, "_designer_custom_yticks", False):
                add(f"ax{idx}.set_yticks({self._format_float_list(ax.get_yticks())})")

            data_plots = getattr(ax, "_designer_data_plots", [])
            if data_plots:
                for plot in data_plots:
                    kind = plot.get("kind")
                    label = plot.get("field")
                    color = plot.get("color")
                    if kind == "line":
                        parts = [
                            f"color={color!r}" if color is not None else None,
                            "linewidth=2.0",
                            "linestyle='-'",
                        ]
                        parts = [part for part in parts if part is not None]
                        add(
                            f"ax{idx}.plot({self._format_float_list(plot.get('x', []))}, {self._format_float_list(plot.get('y', []))}, "
                            f"label={label!r}, {', '.join(parts)})"
                        )
                    elif kind == "scatter":
                        if color is None:
                            add(
                                f"ax{idx}.scatter({self._format_float_list(plot.get('x', []))}, {self._format_float_list(plot.get('y', []))}, "
                                f"s=24, label={label!r})"
                            )
                        else:
                            add(
                                f"ax{idx}.scatter({self._format_float_list(plot.get('x', []))}, {self._format_float_list(plot.get('y', []))}, "
                                f"s=24, label={label!r}, color={color!r})"
                            )
                    elif kind == "bar":
                        if color is None:
                            add(
                                f"ax{idx}.bar({self._format_float_list(plot.get('x', []))}, {self._format_float_list(plot.get('y', []))}, "
                                f"width={float(plot.get('width', 0.8))!r}, label={label!r})"
                            )
                        else:
                            add(
                                f"ax{idx}.bar({self._format_float_list(plot.get('x', []))}, {self._format_float_list(plot.get('y', []))}, "
                                f"width={float(plot.get('width', 0.8))!r}, label={label!r}, color={color!r})"
                            )
                    elif kind == "hist":
                        if color is None:
                            add(
                                f"ax{idx}.hist({self._format_float_list(plot.get('values', []))}, bins={plot.get('bins', 'auto')!r}, "
                                f"alpha={float(plot.get('alpha', 0.55))!r}, label={label!r})"
                            )
                        else:
                            add(
                                f"ax{idx}.hist({self._format_float_list(plot.get('values', []))}, bins={plot.get('bins', 'auto')!r}, "
                                f"alpha={float(plot.get('alpha', 0.55))!r}, label={label!r}, color={color!r})"
                            )
                if len(data_plots) > 1 or any(plot.get("kind") in {"line", "scatter", "bar", "hist"} for plot in data_plots):
                    add(f"ax{idx}.legend()")
            else:
                for line in ax.lines:
                    xdata = list(line.get_xdata())
                    ydata = list(line.get_ydata())
                    add(
                        f"ax{idx}.plot({self._format_float_list(xdata)}, {self._format_float_list(ydata)}, "
                        f"color={line.get_color()!r}, linewidth={float(line.get_linewidth())!r}, "
                        f"linestyle={line.get_linestyle()!r})"
                    )

                for collection in ax.collections:
                    offsets = getattr(collection, "get_offsets", lambda: [])()
                    if offsets is None or len(offsets) == 0:
                        continue
                    xs = [float(point[0]) for point in offsets]
                    ys = [float(point[1]) for point in offsets]
                    facecolors = collection.get_facecolors()
                    color = None
                    if len(facecolors) > 0:
                        color = self._rgba_tuple(facecolors[0])
                    sizes = collection.get_sizes()
                    size = float(sizes[0]) if len(sizes) > 0 else 18.0
                    if color is None:
                        add(f"ax{idx}.scatter({self._format_float_list(xs)}, {self._format_float_list(ys)}, s={size!r})")
                    else:
                        add(
                            f"ax{idx}.scatter({self._format_float_list(xs)}, {self._format_float_list(ys)}, "
                            f"s={size!r}, color={color!r})"
                        )

                legend_labels = [
                    artist.get_label()
                    for artist in list(ax.lines) + list(ax.collections)
                    if artist.get_label() and not str(artist.get_label()).startswith("_")
                ]
                if legend_labels:
                    add(f"ax{idx}.legend()")

            for record in self._get_arrow_records(ax):
                start = tuple(float(v) for v in record["start"])
                end = tuple(float(v) for v in record["end"])
                label = record["label"]
                add(
                    f"arrow = FancyArrowPatch({start!r}, {end!r}, arrowstyle='->', mutation_scale=14, "
                    "linewidth=1.6, color='#111827')"
                )
                add(f"ax{idx}.add_patch(arrow)")
                add(
                    f"ax{idx}.text({(start[0] + end[0]) / 2!r}, {(start[1] + end[1]) / 2!r}, {label.get_text()!r}, "
                    "fontsize=10, ha='center', va='bottom', color='#111827', "
                    "bbox=dict(boxstyle='round,pad=0.2', fc='white', ec='#6b7280', alpha=0.92))"
                )

            for text in ax.texts:
                if id(text) in arrow_label_ids:
                    continue
                x, y = text.get_position()
                bbox = text.get_bbox_patch()
                parts = [
                    f"fontsize={float(text.get_fontsize())!r}",
                    f"ha={text.get_ha()!r}",
                    f"va={text.get_va()!r}",
                    f"color={text.get_color()!r}",
                ]
                if bbox is not None:
                    facecolor = self._rgba_tuple(bbox.get_facecolor())
                    edgecolor = self._rgba_tuple(bbox.get_edgecolor())
                    alpha = bbox.get_alpha()
                    parts.append(
                        "bbox=dict(" 
                        f"boxstyle='round,pad=0.25', fc={facecolor!r}, ec={edgecolor!r}, alpha={float(alpha) if alpha is not None else 1.0!r})"
                    )
                add(f"ax{idx}.text({float(x)!r}, {float(y)!r}, {text.get_text()!r}, {', '.join(parts)})")

            add("")

        add("if __name__ == '__main__':")
        add("    plt.show()")
        return "\n".join(lines)

    def _place_annotation(self, ax, x: float, y: float, text: str) -> None:
        ax.text(
            x,
            y,
            text,
            fontsize=10,
            color="#111827",
            bbox=dict(boxstyle="round,pad=0.25", fc="#fff7cc", ec="#a67c00", alpha=0.96),
            picker=True,
        )

    def _place_arrow(self, ax, start: tuple[float, float], end: tuple[float, float], text: str) -> None:
        arrow = FancyArrowPatch(
            start,
            end,
            arrowstyle="->",
            mutation_scale=14,
            linewidth=1.6,
            color="#111827",
            picker=5,
        )
        ax.add_patch(arrow)
        label = ax.text(
            (start[0] + end[0]) / 2,
            (start[1] + end[1]) / 2,
            text,
            fontsize=10,
            ha="center",
            va="bottom",
            color="#111827",
            bbox=dict(boxstyle="round,pad=0.2", fc="white", ec="#6b7280", alpha=0.92),
            picker=True,
        )
        if not hasattr(ax, "_designer_arrows"):
            ax._designer_arrows = []
        ax._designer_arrows.append({"arrow": arrow, "label": label, "start": start, "end": end})

    def on_canvas_click(self, event) -> None:
        if self.pending_action is None or event.inaxes is None:
            return

        ax = event.inaxes
        x, y = event.xdata, event.ydata

        if self.pending_action.kind == "annotation":
            self._place_annotation(ax, x, y, self.pending_action.text)
            self.pending_action = None
            self._set_status("Annotation placed. You can drag it.")
            self.refresh_all()
            return

        if self.pending_action.kind == "arrow":
            if self.pending_action.start is None:
                self.pending_action.start = (x, y)
                self._set_status("Arrow start captured. Click the arrow end point.")
                return
            self._place_arrow(ax, self.pending_action.start, (x, y), self.pending_action.text)
            self.pending_action = None
            self._set_status("Arrow placed. You can drag the arrow or its label.")
            self.refresh_all()

    def _get_arrow_records(self, ax) -> list[dict]:
        return getattr(ax, "_designer_arrows", [])

    def _find_arrow_record(self, artist):
        for item in self.axes_items:
            ax = item["ax"]
            for record in self._get_arrow_records(ax):
                if record["arrow"] is artist or record["label"] is artist:
                    return ax, record
        return None, None

    def on_pick(self, event) -> None:
        artist = event.artist
        ax, record = self._find_arrow_record(artist)
        if record is not None and ax is not None:
            mouse_x = event.mouseevent.xdata
            mouse_y = event.mouseevent.ydata
            if mouse_x is None or mouse_y is None:
                return
            self.drag_state = DragState(
                kind="arrow",
                axes=ax,
                artist=artist,
                start_mouse=(mouse_x, mouse_y),
                start_data=(0.0, 0.0),
                payload={"record": record, "start": record["start"], "end": record["end"]},
            )
            self._set_status("Dragging arrow. Release to place.")
            return

        if hasattr(artist, "axes") and getattr(artist, "axes", None) is not None:
            if getattr(artist, "get_text", None) is not None and artist in artist.axes.texts:
                mouse_x = event.mouseevent.xdata
                mouse_y = event.mouseevent.ydata
                if mouse_x is None or mouse_y is None:
                    return
                self.drag_state = DragState(
                    kind="text",
                    axes=artist.axes,
                    artist=artist,
                    start_mouse=(mouse_x, mouse_y),
                    start_data=artist.get_position(),
                    payload={},
                )
                self._set_status("Dragging text item. Release to place.")

    def on_motion(self, event) -> None:
        if self.drag_state is None or event.inaxes is None:
            return
        if event.xdata is None or event.ydata is None:
            return
        if event.inaxes != self.drag_state.axes:
            return

        dx = event.xdata - self.drag_state.start_mouse[0]
        dy = event.ydata - self.drag_state.start_mouse[1]

        if self.drag_state.kind == "text":
            artist = self.drag_state.artist
            x0, y0 = self.drag_state.start_data
            artist.set_position((x0 + dx, y0 + dy))
            self.canvas.draw_idle()
            return

        if self.drag_state.kind == "arrow":
            record = self.drag_state.payload["record"]
            start = self.drag_state.payload["start"]
            end = self.drag_state.payload["end"]
            new_start = (start[0] + dx, start[1] + dy)
            new_end = (end[0] + dx, end[1] + dy)
            record["start"] = new_start
            record["end"] = new_end
            record["arrow"].set_positions(new_start, new_end)
            record["label"].set_position(((new_start[0] + new_end[0]) / 2, (new_start[1] + new_end[1]) / 2))
            self.canvas.draw_idle()

    def on_release(self, _event) -> None:
        if self.drag_state is not None:
            self.drag_state = None
            self._set_status("Drag complete.")
            self.canvas.draw_idle()


def main() -> None:
    root = tk.Tk()
    style = ttk.Style(root)
    if "clam" in style.theme_names():
        style.theme_use("clam")
    MatplotlibDesignerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()

